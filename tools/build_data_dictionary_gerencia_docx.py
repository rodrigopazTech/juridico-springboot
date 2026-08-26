import json
import os
import sys
from datetime import datetime
from pathlib import Path

from docx import Document
from docx.enum.section import WD_ORIENTATION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

TOOLS_DIR = Path(__file__).resolve().parent
BASE_DIR = TOOLS_DIR.parent
sys.path.insert(0, str(TOOLS_DIR))

from build_data_dictionary_docx import (  # noqa: E402
    MODULES,
    TABLE_OVERVIEWS,
    describe_column,
    draw_er,
    fk_lookup,
    load_schema,
    module_for,
    table_overview,
)

OUT_DIR = BASE_DIR / "documentacion"
DOCX_PATH = OUT_DIR / "Diccionario_de_datos_Formato_Gerencia.docx"

GREEN = "009353"
DARK_GREEN = RGBColor(0, 104, 65)
LIGHT_GREEN = "BEE3D3"
GRAY = "DDDDDD"
SOFT_GRAY = "F3F3F3"


def shade_cell(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    tc_pr.append(shd)


def set_cell_margins(cell, top=80, start=80, bottom=80, end=80):
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for name, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{name}"))
        if node is None:
            node = OxmlElement(f"w:{name}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def cell_text(cell, value, bold=False, size=7.6, color=None, align=None):
    cell.text = ""
    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(0)
    if align:
        p.alignment = align
    run = p.add_run("" if value is None else str(value))
    run.bold = bold
    run.font.name = "Calibri"
    run.font.size = Pt(size)
    if color:
        run.font.color.rgb = color


def repeat_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def set_table_widths(table, widths):
    for row in table.rows:
        for i, width in enumerate(widths):
            if i < len(row.cells):
                row.cells[i].width = Inches(width)


def set_table_fixed_layout(table):
    tbl_pr = table._tbl.tblPr
    layout = tbl_pr.first_child_found_in("w:tblLayout")
    if layout is None:
        layout = OxmlElement("w:tblLayout")
        tbl_pr.append(layout)
    layout.set(qn("w:type"), "fixed")


def configure_doc(doc):
    section = doc.sections[0]
    section.orientation = WD_ORIENTATION.LANDSCAPE
    section.page_width = Inches(11)
    section.page_height = Inches(8.5)
    section.top_margin = Inches(0.55)
    section.bottom_margin = Inches(0.55)
    section.left_margin = Inches(0.45)
    section.right_margin = Inches(0.45)

    normal = doc.styles["Normal"]
    normal.font.name = "Calibri"
    normal.font.size = Pt(10)
    normal.paragraph_format.space_after = Pt(5)
    normal.paragraph_format.line_spacing = 1.05

    for name, size, before, after in [
        ("Heading 1", 16, 10, 7),
        ("Heading 2", 12, 8, 4),
        ("Heading 3", 10.5, 6, 3),
    ]:
        style = doc.styles[name]
        style.font.name = "Calibri"
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = DARK_GREEN
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)


def add_cover(doc, schema):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(70)
    r = p.add_run("DICCIONARIO DE DATOS")
    r.bold = True
    r.font.name = "Calibri"
    r.font.size = Pt(30)
    r.font.color.rgb = DARK_GREEN

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("Sistema Juridico")
    r.bold = True
    r.font.size = Pt(18)

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("Formato tecnico de gerencia")
    r.font.size = Pt(13)
    r.font.color.rgb = RGBColor(80, 80, 80)

    tbl = doc.add_table(rows=4, cols=2)
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    tbl.style = "Table Grid"
    rows = [
        ("Base de datos", "sistema_juridico"),
        ("Fecha de generacion", datetime.now().strftime("%Y-%m-%d %H:%M")),
        ("Tablas documentadas", str(len(schema))),
        ("Relaciones declaradas", str(sum(len(t["foreign_keys"]) for t in schema))),
    ]
    for i, (label, value) in enumerate(rows):
        for c in tbl.rows[i].cells:
            set_cell_margins(c, 120, 120, 120, 120)
            c.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        shade_cell(tbl.rows[i].cells[0], GREEN)
        cell_text(tbl.rows[i].cells[0], label, True, 10, RGBColor(255, 255, 255))
        cell_text(tbl.rows[i].cells[1], value, False, 10)
    set_table_widths(tbl, [2.3, 4.2])

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(18)
    r = p.add_run("Documento generado a partir de la estructura actual de PostgreSQL y alineado al formato de diccionario utilizado por la gerencia.")
    r.font.size = Pt(10)
    r.font.color.rgb = RGBColor(90, 90, 90)
    doc.add_page_break()


def add_relevant_data(doc, schema):
    operational = [t for t in schema if not t["table"].startswith("bk_")]
    backups = [t for t in schema if t["table"].startswith("bk_")]
    columns = sum(len(t["columns"]) for t in schema)
    fks = sum(len(t["foreign_keys"]) for t in schema)
    pks = sum(1 for t in schema if t.get("primary_key"))
    uniques = sum(len(t.get("unique_constraints", [])) for t in schema)
    modules = sorted({module_for(t["table"]) for t in schema})

    doc.add_heading("Datos relevantes de la base de datos", level=1)
    doc.add_paragraph(
        "El presente Diccionario de Datos es un documento tecnico que describe de manera estructurada y detallada los elementos que conforman la base de datos del sistema. "
        "Su finalidad es proporcionar una guia clara sobre la organizacion, significado y caracteristicas de cada uno de los datos utilizados dentro del sistema."
    )

    tbl = doc.add_table(rows=1, cols=2)
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    tbl.style = "Table Grid"
    metrics = [
        ("Total de tablas", len(schema)),
        ("Tablas operativas", len(operational)),
        ("Tablas de respaldo historico", len(backups)),
        ("Columnas documentadas", columns),
        ("Relaciones FK", fks),
        ("Tablas con llave primaria", pks),
        ("Restricciones UNIQUE", uniques),
        ("Modulos identificados", ", ".join(modules)),
    ]
    for i, (label, value) in enumerate(metrics):
        row = tbl.rows[0] if i == 0 else tbl.add_row()
        for c in row.cells:
            set_cell_margins(c, 110, 120, 110, 120)
        shade_cell(row.cells[0], GREEN if i == 0 else SOFT_GRAY)
        cell_text(row.cells[0], label, True, 9, RGBColor(255, 255, 255) if i == 0 else RGBColor(0, 0, 0))
        cell_text(row.cells[1], value, False, 9)
    set_table_widths(tbl, [2.2, 7.2])

    doc.add_heading("Criterios de lectura", level=2)
    crit = [
        ("Nombre Tabla", "Nombre fisico de la tabla en PostgreSQL."),
        ("Nombre Columna", "Nombre fisico del campo."),
        ("Tipo de Dato", "Tipo SQL declarado en la base."),
        ("Longitud Max", "Longitud maxima declarada cuando el tipo lo permite."),
        ("Requisito", "Indica si el campo permite NULL o es obligatorio."),
        ("Tipo de Llave (PK/FK)", "Indica llave primaria, llave foranea o ambas cuando corresponde."),
        ("Descripcion Funcional", "Uso de negocio o tecnico del campo dentro del sistema."),
    ]
    tbl = doc.add_table(rows=1, cols=2)
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    tbl.style = "Table Grid"
    for i, (label, value) in enumerate(crit):
        row = tbl.rows[0] if i == 0 else tbl.add_row()
        for c in row.cells:
            set_cell_margins(c, 90, 100, 90, 100)
        shade_cell(row.cells[0], GREEN if i == 0 else SOFT_GRAY)
        cell_text(row.cells[0], label, True, 8.8, RGBColor(255, 255, 255) if i == 0 else RGBColor(0, 0, 0))
        cell_text(row.cells[1], value, False, 8.8)
    set_table_widths(tbl, [2.1, 7.3])
    doc.add_page_break()


def type_and_length(col):
    dtype = col.get("data_type") or col.get("type") or ""
    max_len = col.get("max_length")
    if max_len is None and "(" in dtype and dtype.endswith(")"):
        base, length = dtype.rsplit("(", 1)
        return base.strip(), length[:-1]
    return dtype, "" if max_len is None else str(max_len)


def key_type(table, col, fk_map):
    keys = []
    if col["column"] in set(table.get("primary_key", [])):
        keys.append("PK")
    if fk_map.get((table["table"], col["column"])):
        keys.append("FK")
    return " / ".join(keys)


def add_dictionary_table(doc, table, fk_map):
    headers = [
        "Nombre Tabla",
        "Nombre Columna",
        "Tipo de Dato",
        "Longitud Max",
        "Requisito",
        "Tipo de Llave (PK/FK)",
        "Descripcion Funcional",
    ]
    widths = [1.55, 1.55, 1.1, 0.75, 1.25, 1.0, 3.1]
    tbl = doc.add_table(rows=1, cols=len(headers))
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    tbl.style = "Table Grid"
    set_table_fixed_layout(tbl)
    repeat_header(tbl.rows[0])
    for i, header in enumerate(headers):
        cell = tbl.rows[0].cells[i]
        shade_cell(cell, GREEN)
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        set_cell_margins(cell)
        cell_text(cell, header, True, 7.4, RGBColor(255, 255, 255), WD_ALIGN_PARAGRAPH.CENTER)

    for idx, col in enumerate(table["columns"]):
        row = tbl.add_row()
        fill = GRAY if idx % 2 == 1 else None
        if idx == len(table["columns"]) - 1:
            fill = LIGHT_GREEN
        dtype, max_len = type_and_length(col)
        vals = [
            table["table"],
            col["column"],
            dtype,
            max_len,
            "Obligatorio (NOT NULL)" if col["nullable"] == "NO" else "Opcional (NULL)",
            key_type(table, col, fk_map),
            describe_column(table["table"], col, fk_map),
        ]
        for i, val in enumerate(vals):
            cell = row.cells[i]
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            set_cell_margins(cell)
            if fill:
                shade_cell(cell, fill)
            cell_text(cell, val, False, 6.9)
    set_table_widths(tbl, widths)


def add_table_sections(doc, schema):
    fk_map = fk_lookup(schema)
    doc.add_heading("Diccionario de datos", level=1)
    for table in sorted(schema, key=lambda x: (module_for(x["table"]), x["table"])):
        doc.add_heading(f"Tabla:{table['table']}", level=2)
        p = doc.add_paragraph(table_overview(table["table"]))
        p.paragraph_format.space_after = Pt(5)
        add_dictionary_table(doc, table, fk_map)
        doc.add_paragraph()


def add_relationship_map(doc, schema):
    er_path = draw_er(schema)
    doc.add_page_break()
    doc.add_heading("Mapa relacional", level=1)
    doc.add_paragraph(
        "El siguiente mapa relacional muestra la agrupacion de tablas por modulo, sus llaves primarias y las principales relaciones foraneas declaradas en la base de datos actual."
    )
    doc.add_picture(er_path, width=Inches(9.8))
    doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER


def add_footer(doc):
    for section in doc.sections:
        footer = section.footer.paragraphs[0]
        footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = footer.add_run("Diccionario de Datos - Sistema Juridico")
        run.font.size = Pt(8)
        run.font.color.rgb = RGBColor(90, 90, 90)


def build():
    OUT_DIR.mkdir(parents=True, exist_ok=True)
    schema = load_schema()
    doc = Document()
    configure_doc(doc)
    add_cover(doc, schema)
    add_relevant_data(doc, schema)
    add_table_sections(doc, schema)
    add_relationship_map(doc, schema)
    add_footer(doc)
    doc.core_properties.title = "Diccionario de datos - Formato Gerencia"
    doc.core_properties.subject = "Diccionario de datos de sistema_juridico"
    doc.core_properties.author = "Codex"
    doc.save(DOCX_PATH)
    return DOCX_PATH


if __name__ == "__main__":
    print(build())
