from copy import copy
from pathlib import Path

import openpyxl


ROOT = Path(__file__).resolve().parents[1]
SOURCE = ROOT / "outputs" / "vulnerabilidades_siaj" / "SIAJ_Vulnerabilidades_Formato_Seguridad_2026-08-18.xlsx"
OUTPUT_DIR = ROOT / "outputs" / "vulnerabilidades_siaj"
OUTPUT = OUTPUT_DIR / "SIAJ_Vulnerabilidades_Atendidas_Formato_Seguridad_2026-08-18.xlsx"


def normalize(value):
    return str(value or "").strip().upper()


def copy_cell_style(src, dst):
    if src.has_style:
        dst._style = copy(src._style)
    if src.number_format:
        dst.number_format = src.number_format
    if src.alignment:
        dst.alignment = copy(src.alignment)
    if src.protection:
        dst.protection = copy(src.protection)


def copy_row_style(ws, source_row, target_row, max_col=7):
    for col in range(1, max_col + 1):
        copy_cell_style(ws.cell(source_row, col), ws.cell(target_row, col))
    ws.row_dimensions[target_row].height = ws.row_dimensions[source_row].height


def status_column(ws):
    for col in range(1, ws.max_column + 1):
        if normalize(ws.cell(4, col).value) == "ESTADO":
            return col
    return None


def set_table_range(ws):
    max_col = max(6, min(ws.max_column, 7))
    max_row = max(ws.max_row, 6)
    for table in ws.tables.values():
        table.ref = f"A1:{openpyxl.utils.get_column_letter(max_col)}{max_row}"


def main():
    OUTPUT_DIR.mkdir(parents=True, exist_ok=True)
    wb = openpyxl.load_workbook(SOURCE)

    totals = {}
    for ws in wb.worksheets:
        status_col = status_column(ws)
        if not status_col:
            totals[ws.title] = 0
            continue

        resolved_rows = []
        for row in range(5, ws.max_row + 1):
            vulnerability = ws.cell(row, 1).value
            status = normalize(ws.cell(row, status_col).value)
            if vulnerability and status == "EN PROCESO":
                values = [ws.cell(row, col).value for col in range(1, ws.max_column + 1)]
                values[status_col - 1] = "RESUELTO"
                resolved_rows.append(values)

        if ws.max_row >= 5:
            ws.delete_rows(5, ws.max_row - 4)

        if not resolved_rows:
            copy_row_style(ws, 5, 5, max_col=max(ws.max_column, 7))
            for col in range(1, ws.max_column + 1):
                ws.cell(5, col).value = None
            set_table_range(ws)
            totals[ws.title] = 0
            continue

        for offset, values in enumerate(resolved_rows, start=5):
            if offset > ws.max_row:
                ws.append([None] * ws.max_column)
            copy_row_style(ws, 5, offset, max_col=max(ws.max_column, 7))
            for col, value in enumerate(values, start=1):
                ws.cell(offset, col).value = value

        set_table_range(ws)
        totals[ws.title] = len(resolved_rows)

    wb.save(OUTPUT)
    print(OUTPUT)
    for sheet, count in totals.items():
        print(f"{sheet}: {count}")


if __name__ == "__main__":
    main()
