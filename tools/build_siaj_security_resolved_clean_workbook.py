from copy import copy
from pathlib import Path

import openpyxl
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.worksheet.table import Table, TableStyleInfo


ROOT = Path(__file__).resolve().parents[1]
SOURCE = ROOT / "outputs" / "vulnerabilidades_siaj" / "SIAJ_Plan_Vulnerabilidades_2026_2028.xlsx"
OUTPUT_DIR = ROOT / "outputs" / "vulnerabilidades_siaj"
OUTPUT = OUTPUT_DIR / "SIAJ_Vulnerabilidades_Atendidas_Formato_Seguridad_Limpio_2026-08-18.xlsx"

SEVERITIES = ["CRITICAL", "HIGH", "MEDIUM", "LOW"]
HEADERS = ["Vulnerabilidad", "Package Name", "Categoria", "PáginaDF", "ESTADO"]

COLORS = {
    "CRITICAL": "C00000",
    "HIGH": "70AD47",
    "MEDIUM": "FFC000",
    "LOW": "9EADCC",
}


def normalize(value):
    return str(value or "").strip().upper()


def copy_cell_style(src, dst):
    if src.has_style:
        dst._style = copy(src._style)
    if src.number_format:
        dst.number_format = src.number_format
    if src.alignment:
        dst.alignment = copy(src.alignment)
    if src.protection:
        dst.protection = copy(src.protection)


def source_headers(ws):
    return {normalize(ws.cell(4, col).value): col for col in range(1, ws.max_column + 1)}


def col(indexes, *names):
    for name in names:
        key = normalize(name)
        if key in indexes:
            return indexes[key]
    raise KeyError(names)


def read_resolved_candidates():
    wb = openpyxl.load_workbook(SOURCE, data_only=True)
    ws = wb["COMPENDIO"]
    indexes = source_headers(ws)
    cols = {
        "severity": col(indexes, "GRAVEDAD"),
        "vulnerability": col(indexes, "VULNERABILIDAD"),
        "package": col(indexes, "PACKAGE", "PACKAGE NAME"),
        "category": col(indexes, "CATEGORIA / CWE", "CATEGORIA", "CATEGORY"),
        "page": col(indexes, "PAGINA", "PÁGINA", "PAGINADF"),
        "status": col(indexes, "ESTATUS", "ESTADO"),
    }

    rows = {severity: [] for severity in SEVERITIES}
    for row in range(5, ws.max_row + 1):
        status = normalize(ws.cell(row, cols["status"]).value)
        if status != "EN PROCESO":
            continue
        severity = normalize(ws.cell(row, cols["severity"]).value)
        if severity not in rows:
            rows[severity] = []
        rows[severity].append([
            ws.cell(row, cols["vulnerability"]).value,
            ws.cell(row, cols["package"]).value,
            ws.cell(row, cols["category"]).value,
            ws.cell(row, cols["page"]).value,
            "RESUELTO",
        ])
    return rows


def style_sheet(ws, severity, rows):
    color = COLORS.get(severity, "70AD47")
    light_fill = {
        "CRITICAL": "F4CCCC",
        "HIGH": "E2F0D9",
        "MEDIUM": "FFF2CC",
        "LOW": "D9E2F3",
    }.get(severity, "E2F0D9")

    ws.sheet_view.showGridLines = False
    ws.freeze_panes = "A5"

    for col, width in zip("ABCDE", [58, 42, 22, 14, 16]):
        ws.column_dimensions[col].width = width

    ws.merge_cells("A1:E1")
    ws["A1"] = f"VULNERABILIDADES ATENDIDAS - {severity}"
    ws["A1"].fill = PatternFill("solid", fgColor=color)
    ws["A1"].font = Font(color="FFFFFF", bold=True, size=13)
    ws["A1"].alignment = Alignment(horizontal="center", vertical="center")
    ws.row_dimensions[1].height = 24

    ws.merge_cells("A2:E2")
    ws["A2"] = "Sistema Integral de Agenda Jurídica (SIAJ)"
    ws["A2"].fill = PatternFill("solid", fgColor=light_fill)
    ws["A2"].font = Font(bold=True)
    ws["A2"].alignment = Alignment(horizontal="center")

    ws.append([])
    ws.append(HEADERS)

    header_fill = PatternFill("solid", fgColor="1F4E78")
    header_font = Font(color="FFFFFF", bold=True)
    thin = Side(style="thin", color="D9E2F3")
    border = Border(top=thin, left=thin, right=thin, bottom=thin)

    for cell in ws[4]:
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", vertical="center")
        cell.border = border

    if rows:
        for item in rows:
            ws.append(item)
    else:
        ws.append(["Sin vulnerabilidades atendidas en esta severidad", "", "", "", ""])

    for row in range(5, ws.max_row + 1):
        for col in range(1, 6):
            cell = ws.cell(row, col)
            cell.border = border
            cell.alignment = Alignment(
                horizontal="center" if col in (3, 4, 5) else "left",
                vertical="center",
                wrap_text=True,
            )
            if row % 2 == 0:
                cell.fill = PatternFill("solid", fgColor="F8FBFD")

        status_cell = ws.cell(row, 5)
        if status_cell.value == "RESUELTO":
            status_cell.fill = PatternFill("solid", fgColor="A9D18E")
            status_cell.font = Font(color="000000", bold=True)

    table_ref = f"A4:E{ws.max_row}"
    table = Table(displayName=f"Tabla{severity.title()}", ref=table_ref)
    table.tableStyleInfo = TableStyleInfo(
        name="TableStyleMedium2",
        showFirstColumn=False,
        showLastColumn=False,
        showRowStripes=False,
        showColumnStripes=False,
    )
    ws.add_table(table)


def main():
    OUTPUT_DIR.mkdir(parents=True, exist_ok=True)
    rows_by_severity = read_resolved_candidates()

    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    for severity in SEVERITIES:
        ws = wb.create_sheet(severity)
        style_sheet(ws, severity, rows_by_severity.get(severity, []))

    wb.save(OUTPUT)
    print(OUTPUT)
    for severity in SEVERITIES:
        print(f"{severity}: {len(rows_by_severity.get(severity, []))}")


if __name__ == "__main__":
    main()
