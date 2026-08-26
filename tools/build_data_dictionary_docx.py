import json
import math
import os
import re
import textwrap
from collections import defaultdict
from datetime import datetime

from PIL import Image, ImageDraw, ImageFont
from docx import Document
from docx.enum.section import WD_ORIENTATION
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


BASE_DIR = os.path.abspath(os.path.join(os.path.dirname(__file__), ".."))
SCHEMA_JSON = os.path.join(BASE_DIR, "db_schema_actual.json")
OUT_DIR = os.path.join(BASE_DIR, "documentacion")
ER_IMAGE = os.path.join(OUT_DIR, "diccionario_datos_er.png")
DOCX_PATH = os.path.join(OUT_DIR, "Diccionario de datos.docx")


TABLE_OVERVIEWS = {
    "usuarios": "Maestro de usuarios del sistema. Define identidad, rol, datos de contacto, autenticacion local y datos de enlace con identidad Microsoft historica.",
    "usuario_gerencias": "Tabla puente que asigna usuarios a una o varias gerencias para filtrar permisos, bandejas y consultas operativas.",
    "usuario_materias": "Tabla puente que asigna usuarios a materias juridicas; se usa para alcance operativo, responsables y notificaciones.",
    "usuario_permisos_adicionales": "Permisos granulares adicionales a los derivados del rol principal del usuario.",
    "gerencias": "Catalogo de gerencias juridicas que organizan expedientes, materias, organos, usuarios y configuraciones.",
    "materias": "Catalogo de materias juridicas por gerencia, por ejemplo Laboral, Penal, Civil, Mercantil, Fiscal, Transparencia o Amparo.",
    "cat_estados": "Catalogo de estados o estatus configurables, asociados a gerencia y usuario creador cuando aplica.",
    "cat_prioridad": "Catalogo de prioridades visuales y operativas para recordatorios y notificaciones.",
    "cat_tipo_expediente": "Catalogo de tipos de expediente usados para clasificar asuntos juridicos.",
    "organos_jurisdiccionales": "Catalogo de organos jurisdiccionales o autoridades ante las que se tramitan asuntos.",
    "tipos_audiencia": "Catalogo de tipos de audiencia disponibles por gerencia.",
    "expedientes": "Entidad central del sistema. Registra cada asunto juridico, su numero, partes, materia, gerencia, etapa, fechas y responsable.",
    "documentos_expediente": "Relacion historica o auxiliar entre documentos y expedientes, usada por cargas o asociaciones documentales.",
    "documentos": "Repositorio logico de archivos cargados en expedientes o carpetas, con metadatos de almacenamiento, eliminacion y auditoria.",
    "carpetas": "Estructura de carpetas por expediente para organizar documentos y mantener jerarquia documental.",
    "actividad_expedientes": "Bitacora funcional de eventos visibles de un expediente, como cambios, cargas o acciones de usuario.",
    "bitacora_datos": "Auditoria detallada de cambios de campos en expedientes, con valor anterior y nuevo.",
    "colaboradores_expediente": "Usuarios invitados o colaboradores temporales de un expediente, con vigencia y nivel de permiso.",
    "audiencias": "Registro de audiencias programadas o gestionadas dentro de un expediente.",
    "audiencia_contextos": "Snapshot de contexto juridico usado para audiencias, especialmente cuando el origen viene de migracion o integraciones.",
    "audiencias_desahogadas": "Registro de audiencias ya desahogadas o sincronizadas como cumplidas.",
    "terminos": "Registro de terminos procesales, fechas limite, responsables, estatus y origen.",
    "termino_contextos": "Snapshot de contexto juridico usado para terminos procesales.",
    "termino_archivos": "Archivos o documentos asociados directamente a un termino procesal.",
    "terminos_presentados": "Registro de presentacion o cumplimiento de terminos procesales.",
    "comentarios": "Comentarios generales asociados de forma polimorfica a expedientes, terminos, audiencias u otras entidades.",
    "notificaciones": "Mensajes de notificacion internos para usuarios, con prioridad, evento y entidad relacionada.",
    "preferencias_notificacion": "Preferencias generales de canales de notificacion por usuario.",
    "preferencias_notificacion_evento": "Preferencias por tipo de evento para activar o desactivar notificaciones especificas.",
    "push_subscriptions": "Suscripciones push web por usuario/dispositivo para notificaciones del navegador.",
    "recordatorios": "Recordatorios personales o de trabajo vinculados a usuario y prioridad.",
    "eventos_calendario": "Eventos de calendario internos asociados a usuarios y, opcionalmente, a entidades del sistema.",
    "actividad_usuarios": "Actividad reciente por usuario para historiales, tableros o auditoria funcional.",
    "auditoria_eventos": "Auditoria tecnica de eventos HTTP, seguridad, acceso y acciones relevantes.",
    "auditoria_sesiones": "Auditoria tecnica de sesiones de usuario, inicio, actividad y cierre.",
    "configuraciones_alerta": "Reglas de alertamiento por nivel, gerencia, materia o tipo de evento.",
    "configuraciones_semaforo": "Configuracion de semaforos de vencimiento y clasificacion temporal.",
    "catalogos_migracion": "Catalogos normalizados o equivalencias usados por procesos de migracion.",
    "lotes_migracion": "Encabezado de lotes de migracion, con responsable, estado y gerencia global.",
    "registros_migracion": "Detalle de registros procesados dentro de cada lote de migracion, incluyendo resultado y errores.",
    "password_reset_tokens": "Tokens de activacion, recuperacion y restablecimiento de contrasena.",
}


MODULES = [
    ("Usuarios y seguridad", ["usuarios", "usuario_", "password_", "preferencias_", "push_", "recordatorios", "notificaciones", "actividad_usuarios", "eventos_calendario"]),
    ("Catalogos", ["gerencias", "materias", "cat_", "organos_", "tipos_audiencia"]),
    ("Expedientes y documentos", ["expedientes", "documentos", "carpetas", "actividad_expedientes", "bitacora_", "colaboradores_"]),
    ("Procesal", ["audiencias", "audiencia_", "terminos", "termino_"]),
    ("Auditoria", ["auditoria_"]),
    ("Configuracion", ["configuraciones_"]),
    ("Migracion y respaldo", ["migracion", "lotes_", "registros_", "bk_"]),
]


def load_schema():
    with open(SCHEMA_JSON, encoding="utf-8-sig") as f:
        return json.load(f)


def module_for(table):
    for name, patterns in MODULES:
        if any(table.startswith(p) or p in table for p in patterns):
            return name
    return "Otros"


def title_case_table(name):
    return name.replace("_", " ").title()


def fk_lookup(schema):
    fks = defaultdict(list)
    for t in schema:
        for fk in t["foreign_keys"]:
            fks[(t["table"], fk["column"])].append(fk)
    return fks


def pk_set(table):
    return set(table.get("primary_key", []))


COMMON_FIELD_DESCRIPTIONS = {
    "id": "Identificador primario interno de la fila.",
    "created_at": "Fecha y hora de creacion del registro.",
    "updated_at": "Fecha y hora de ultima actualizacion del registro.",
    "activo": "Bandera logica que indica si el registro esta disponible para uso operativo.",
    "nombre": "Nombre visible o descripcion corta del registro.",
    "descripcion": "Descripcion amplia del registro o regla.",
    "email": "Correo electronico asociado al usuario o evento.",
    "username": "Nombre de usuario usado para iniciar sesion o identificar la cuenta.",
    "password_hash": "Hash de la contrasena local; no almacena la contrasena en texto plano.",
    "rol": "Rol funcional que determina permisos principales dentro del sistema.",
    "fecha_registro": "Fecha y hora en que se registro la actividad.",
    "fecha_creacion": "Fecha y hora de generacion del registro o token.",
    "fecha_expiracion": "Fecha y hora limite de vigencia.",
    "fecha_uso": "Fecha y hora en que el token o recurso fue utilizado.",
    "usado": "Indica si el token o recurso ya fue consumido.",
    "expirado": "Indica si el token o recurso ya no es vigente.",
    "titulo": "Titulo visible del recordatorio, notificacion, evento o actividad.",
    "mensaje": "Contenido principal del mensaje o notificacion.",
    "leida": "Indica si la notificacion fue marcada como leida.",
    "fecha_leida": "Fecha y hora en que se marco como leida.",
    "prioridad": "Valor textual o categoria de prioridad operativa.",
    "color": "Color usado para representacion visual en interfaz.",
    "codigo": "Codigo funcional o llave de negocio del registro.",
    "tipo": "Categoria o tipo funcional del registro.",
    "estado": "Estado funcional o resultado del proceso.",
    "estatus": "Estatus operativo del expediente, termino, audiencia o proceso.",
    "resultado": "Resultado de una accion, proceso o evento auditado.",
    "detalle_json": "Detalle estructurado en formato JSON para auditoria o diagnostico.",
    "ip": "Direccion IP registrada para auditoria o sesion.",
    "user_agent": "Agente de usuario del navegador o cliente que realizo la accion.",
    "session_id": "Identificador tecnico de sesion HTTP.",
    "ruta": "Ruta HTTP o recurso relacionado con el evento.",
    "metodo_http": "Metodo HTTP usado en la solicitud.",
    "codigo_http": "Codigo de respuesta HTTP asociado al evento.",
    "entidad_tipo": "Tipo de entidad de negocio relacionada; se usa para asociaciones polimorficas.",
    "entidad_id": "Identificador de la entidad relacionada; se usa junto con entidad_tipo.",
    "modulo": "Modulo funcional donde ocurrio el evento.",
    "accion": "Accion realizada o registrada.",
    "token": "Valor de token utilizado para activacion, recuperacion o seguridad.",
}


def describe_column(table_name, col, fk_map):
    name = col["column"]
    if name in COMMON_FIELD_DESCRIPTIONS:
        desc = COMMON_FIELD_DESCRIPTIONS[name]
    elif name.endswith("_id"):
        base = name[:-3].replace("_", " ")
        desc = f"Identificador relacionado con {base}."
    elif name.startswith("fecha_"):
        desc = "Fecha asociada al evento o hito indicado por el nombre del campo."
    elif name.startswith("hora_"):
        desc = "Hora asociada al evento o hito indicado por el nombre del campo."
    elif name.startswith("microsoft_"):
        desc = "Dato historico de integracion con identidad Microsoft/Outlook; conserva referencia o sincronizacion previa."
    elif name.startswith("bk_"):
        desc = "Campo conservado en tabla de respaldo historico."
    elif "archivo" in name or "ruta" in name:
        desc = "Dato de almacenamiento o ubicacion de archivo/documento."
    elif "numero" in name:
        desc = "Numero, folio o identificador de negocio visible para usuarios."
    elif "snapshot" in name:
        desc = "Copia del dato en el momento del evento para preservar contexto historico."
    elif name.startswith("es_") or name.startswith("tiene_") or name.startswith("permite_"):
        desc = "Bandera booleana de control funcional segun el nombre del campo."
    else:
        desc = f"Dato funcional de {title_case_table(table_name).lower()} identificado como {name.replace('_', ' ')}."

    fks = fk_map.get((table_name, name), [])
    if fks:
        targets = ", ".join(f"{fk['foreign_table']}.{fk['foreign_column']}" for fk in fks)
        desc += f" Relaciona con {targets}."
    return desc


def cell_text(cell, text, bold=False, size=8.5, color=None):
    cell.text = ""
    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(0)
    run = p.add_run(str(text) if text is not None else "")
    run.bold = bold
    run.font.size = Pt(size)
    if color:
        run.font.color.rgb = RGBColor(*color)


def shade_cell(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    tc_pr.append(shd)


def set_cell_margins(cell, top=80, start=120, bottom=80, end=120):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for m, v in [("top", top), ("start", start), ("bottom", bottom), ("end", end)]:
        node = tc_mar.find(qn(f"w:{m}"))
        if node is None:
            node = OxmlElement(f"w:{m}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(v))
        node.set(qn("w:type"), "dxa")


def set_repeat_table_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def set_table_widths(table, widths):
    for row in table.rows:
        for idx, width in enumerate(widths):
            if idx < len(row.cells):
                row.cells[idx].width = Inches(width)


def relation_text(table, fk_map):
    parts = []
    for fk in table["foreign_keys"]:
        parts.append(f"{fk['column']} -> {fk['foreign_table']}.{fk['foreign_column']}")
    return "; ".join(parts) if parts else "Sin llaves foraneas declaradas."


def table_keys(table):
    pk = ", ".join(table.get("primary_key", [])) or "Sin PK declarada"
    uniques = table.get("unique_constraints", [])
    uniq = "; ".join(f"{u['constraint']}({u['column']})" for u in uniques) if uniques else "Sin restricciones unique declaradas"
    return pk, uniq


def table_overview(table_name):
    if table_name in TABLE_OVERVIEWS:
        return TABLE_OVERVIEWS[table_name]
    if table_name.startswith("bk_"):
        return "Tabla de respaldo historico creada para conservar datos previos de autorizacion/gerencias. No debe tratarse como tabla operativa principal."
    return f"Tabla de soporte del modulo {module_for(table_name).lower()}."


def draw_er(schema):
    os.makedirs(OUT_DIR, exist_ok=True)
    tables_by_module = defaultdict(list)
    for t in schema:
        tables_by_module[module_for(t["table"])].append(t)

    modules = [m for m, _ in MODULES if m in tables_by_module] + [m for m in sorted(tables_by_module) if m not in {x[0] for x in MODULES}]
    grid_cols = 4
    col_w = 450
    margin = 40
    header_h = 54
    gap_x = 36
    gap_y = 18

    try:
        font = ImageFont.truetype("arial.ttf", 15)
        font_bold = ImageFont.truetype("arialbd.ttf", 15)
        font_module = ImageFont.truetype("arialbd.ttf", 18)
    except Exception:
        font = ImageFont.load_default()
        font_bold = ImageFont.load_default()
        font_module = ImageFont.load_default()

    module_heights = {}
    for m in modules:
        h = header_h + 16
        for t in sorted(tables_by_module[m], key=lambda x: x["table"]):
            rows = 1 + len(t.get("primary_key", [])) + len(t.get("foreign_keys", []))
            h += 42 + rows * 18 + gap_y
        module_heights[m] = h

    row_heights = []
    for start in range(0, len(modules), grid_cols):
        row_heights.append(max(module_heights[m] for m in modules[start:start + grid_cols]))

    width = margin * 2 + grid_cols * col_w + (grid_cols - 1) * gap_x
    height = margin * 2 + sum(row_heights) + (len(row_heights) - 1) * gap_y
    img = Image.new("RGB", (width, height), "white")
    draw = ImageDraw.Draw(img)

    palette = {
        "Usuarios y seguridad": "#E8EEF5",
        "Catalogos": "#EAF4EA",
        "Expedientes y documentos": "#F3EEF8",
        "Procesal": "#F8F1E8",
        "Auditoria": "#F2F4F7",
        "Configuracion": "#EAF3F8",
        "Migracion y respaldo": "#F7EFEF",
        "Otros": "#F4F4F4",
    }
    border = "#8392A5"
    table_positions = {}
    centers = {}

    row_tops = []
    running_y = margin
    for rh in row_heights:
        row_tops.append(running_y)
        running_y += rh + gap_y

    def module_xy(index):
        row = index // grid_cols
        col = index % grid_cols
        return margin + col * (col_w + gap_x), row_tops[row]

    for ci, m in enumerate(modules):
        x, y = module_xy(ci)
        draw.rounded_rectangle([x, y, x + col_w, y + header_h], radius=10, fill=palette.get(m, "#F4F4F4"), outline=border, width=2)
        draw.text((x + 14, y + 16), m, fill="#0B2545", font=font_module)
        y += header_h + 16
        for t in sorted(tables_by_module[m], key=lambda x: x["table"]):
            lines = []
            pks = t.get("primary_key", [])
            if pks:
                lines.append("PK: " + ", ".join(pks))
            for fk in t.get("foreign_keys", [])[:7]:
                lines.append(f"FK: {fk['column']} -> {fk['foreign_table']}")
            if len(t.get("foreign_keys", [])) > 7:
                lines.append(f"... {len(t['foreign_keys']) - 7} FK mas")
            box_h = 42 + max(1, len(lines)) * 18
            x2, y2 = x + col_w, y + box_h
            fill = "#FFFFFF" if not t["table"].startswith("bk_") else "#FFF7F7"
            draw.rounded_rectangle([x, y, x2, y2], radius=6, fill=fill, outline=border, width=1)
            draw.rectangle([x, y, x2, y + 24], fill="#F2F4F7", outline=border)
            draw.text((x + 8, y + 7), t["table"], fill="#0B2545", font=font_bold)
            yy = y + 31
            for line in lines or ["Sin FK declaradas"]:
                draw.text((x + 12, yy), line[:62], fill="#333333", font=font)
                yy += 18
            table_positions[t["table"]] = (x, y, x2, y2)
            centers[t["table"]] = ((x + x2) // 2, (y + y2) // 2)
            y = y2 + gap_y

    # Draw cross-table relationships lightly behind labels using center-to-center lines.
    overlay = Image.new("RGBA", img.size, (255, 255, 255, 0))
    od = ImageDraw.Draw(overlay)
    for t in schema:
        src = t["table"]
        if src not in centers:
            continue
        for fk in t.get("foreign_keys", []):
            dst = fk["foreign_table"]
            if dst in centers and src != dst:
                od.line([centers[src], centers[dst]], fill=(46, 116, 181, 45), width=2)
    img = Image.alpha_composite(img.convert("RGBA"), overlay).convert("RGB")

    # Repaint boxes on top so lines do not obscure text.
    draw = ImageDraw.Draw(img)
    for ci, m in enumerate(modules):
        x, y = module_xy(ci)
        draw.rounded_rectangle([x, y, x + col_w, y + header_h], radius=10, fill=palette.get(m, "#F4F4F4"), outline=border, width=2)
        draw.text((x + 14, y + 16), m, fill="#0B2545", font=font_module)
        y += header_h + 16
        for t in sorted(tables_by_module[m], key=lambda x: x["table"]):
            x1, y1, x2, y2 = table_positions[t["table"]]
            fill = "#FFFFFF" if not t["table"].startswith("bk_") else "#FFF7F7"
            draw.rounded_rectangle([x1, y1, x2, y2], radius=6, fill=fill, outline=border, width=1)
            draw.rectangle([x1, y1, x2, y1 + 24], fill="#F2F4F7", outline=border)
            draw.text((x1 + 8, y1 + 7), t["table"], fill="#0B2545", font=font_bold)
            lines = []
            if t.get("primary_key"):
                lines.append("PK: " + ", ".join(t["primary_key"]))
            for fk in t.get("foreign_keys", [])[:7]:
                lines.append(f"FK: {fk['column']} -> {fk['foreign_table']}")
            if len(t.get("foreign_keys", [])) > 7:
                lines.append(f"... {len(t['foreign_keys']) - 7} FK mas")
            yy = y1 + 31
            for line in lines or ["Sin FK declaradas"]:
                draw.text((x1 + 12, yy), line[:62], fill="#333333", font=font)
                yy += 18

    img.save(ER_IMAGE, quality=95)
    return ER_IMAGE


def add_field_table(doc, table, fk_map):
    headers = ["Campo", "Tipo", "Nulo", "Llave / relacion", "Significado y uso"]
    widths = [1.25, 1.35, 0.55, 1.55, 3.65]
    tbl = doc.add_table(rows=1, cols=len(headers))
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    tbl.style = "Table Grid"
    set_repeat_table_header(tbl.rows[0])
    for i, h in enumerate(headers):
        cell = tbl.rows[0].cells[i]
        shade_cell(cell, "E8EEF5")
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        set_cell_margins(cell)
        cell_text(cell, h, bold=True, size=8.5, color=(11, 37, 69))
    pk = pk_set(table)
    for col in table["columns"]:
        row = tbl.add_row()
        for c in row.cells:
            c.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            set_cell_margins(c)
        relation = []
        if col["column"] in pk:
            relation.append("PK")
        fks = fk_map.get((table["table"], col["column"]), [])
        relation.extend(f"FK -> {fk['foreign_table']}.{fk['foreign_column']}" for fk in fks)
        if not relation:
            relation.append("-")
        vals = [
            col["column"],
            col["type"],
            "Si" if col["nullable"] == "YES" else "No",
            "; ".join(relation),
            describe_column(table["table"], col, fk_map),
        ]
        for i, val in enumerate(vals):
            cell_text(row.cells[i], val, size=7.8)
    set_table_widths(tbl, widths)
    doc.add_paragraph()


def add_summary_table(doc, schema):
    tbl = doc.add_table(rows=1, cols=5)
    tbl.style = "Table Grid"
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    headers = ["Modulo", "Tabla", "Uso principal", "PK", "FK"]
    for i, h in enumerate(headers):
        cell = tbl.rows[0].cells[i]
        shade_cell(cell, "E8EEF5")
        set_cell_margins(cell)
        cell_text(cell, h, True, 8.5, (11, 37, 69))
    for t in sorted(schema, key=lambda x: (module_for(x["table"]), x["table"])):
        row = tbl.add_row()
        vals = [
            module_for(t["table"]),
            t["table"],
            table_overview(t["table"]),
            ", ".join(t.get("primary_key", [])) or "-",
            str(len(t.get("foreign_keys", []))),
        ]
        for i, val in enumerate(vals):
            set_cell_margins(row.cells[i])
            cell_text(row.cells[i], val, size=7.8)
    set_table_widths(tbl, [1.35, 1.55, 3.25, 0.8, 0.45])


def configure_styles(doc):
    section = doc.sections[0]
    section.orientation = WD_ORIENTATION.LANDSCAPE
    section.page_width = Inches(11)
    section.page_height = Inches(8.5)
    for s in doc.sections:
        s.top_margin = Inches(0.5)
        s.bottom_margin = Inches(0.5)
        s.left_margin = Inches(0.5)
        s.right_margin = Inches(0.5)

    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Calibri"
    normal.font.size = Pt(10.5)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.15

    for style_name, size, color, before, after in [
        ("Heading 1", 16, RGBColor(46, 116, 181), 18, 10),
        ("Heading 2", 13, RGBColor(46, 116, 181), 14, 7),
        ("Heading 3", 12, RGBColor(31, 77, 120), 10, 5),
    ]:
        st = styles[style_name]
        st.font.name = "Calibri"
        st.font.size = Pt(size)
        st.font.color.rgb = color
        st.font.bold = True
        st.paragraph_format.space_before = Pt(before)
        st.paragraph_format.space_after = Pt(after)


def build_doc():
    schema = load_schema()
    os.makedirs(OUT_DIR, exist_ok=True)
    er_path = draw_er(schema)
    fk_map = fk_lookup(schema)

    doc = Document()
    configure_styles(doc)

    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = title.add_run("Diccionario de datos")
    run.bold = True
    run.font.size = Pt(24)
    run.font.color.rgb = RGBColor(11, 37, 69)
    subtitle = doc.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = subtitle.add_run("Base de datos actual: sistema_juridico")
    run.font.size = Pt(12)
    run.font.color.rgb = RGBColor(85, 85, 85)
    meta = doc.add_paragraph()
    meta.alignment = WD_ALIGN_PARAGRAPH.CENTER
    meta.add_run(f"Generado el {datetime.now().strftime('%Y-%m-%d %H:%M')} | Tablas: {len(schema)} | Relaciones FK: {sum(len(t['foreign_keys']) for t in schema)}")

    doc.add_heading("Esquema ER", level=1)
    p = doc.add_paragraph("El diagrama muestra las tablas agrupadas por modulo. Cada caja incluye su llave primaria y las principales llaves foraneas; el detalle completo de cada campo se documenta en las secciones posteriores.")
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    doc.add_picture(er_path, width=Inches(9.8))
    doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER

    doc.add_page_break()
    doc.add_heading("Resumen de tablas", level=1)
    add_summary_table(doc, schema)

    doc.add_page_break()
    doc.add_heading("Diccionario tabla por tabla", level=1)
    for t in sorted(schema, key=lambda x: (module_for(x["table"]), x["table"])):
        doc.add_heading(t["table"], level=2)
        doc.add_paragraph(table_overview(t["table"]))
        pk, uniques = table_keys(t)
        details = doc.add_table(rows=3, cols=2)
        details.style = "Table Grid"
        details.alignment = WD_TABLE_ALIGNMENT.CENTER
        rows = [
            ("Modulo", module_for(t["table"])),
            ("Llave primaria", pk),
            ("Relaciones / restricciones", f"FK: {relation_text(t, fk_map)} | Unique: {uniques}"),
        ]
        for r, (label, value) in enumerate(rows):
            shade_cell(details.rows[r].cells[0], "F2F4F7")
            for c in details.rows[r].cells:
                set_cell_margins(c)
            cell_text(details.rows[r].cells[0], label, True, 8.5, (11, 37, 69))
            cell_text(details.rows[r].cells[1], value, False, 8.2)
        set_table_widths(details, [1.4, 5.8])
        doc.add_paragraph()
        add_field_table(doc, t, fk_map)

    doc.core_properties.title = "Diccionario de datos"
    doc.core_properties.subject = "Diccionario de datos y esquema ER de sistema_juridico"
    doc.core_properties.author = "Codex"
    doc.save(DOCX_PATH)
    return DOCX_PATH


if __name__ == "__main__":
    print(build_doc())
