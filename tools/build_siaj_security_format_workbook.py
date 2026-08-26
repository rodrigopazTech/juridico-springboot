from copy import copy
from pathlib import Path

import openpyxl


ROOT = Path(__file__).resolve().parents[1]
TEMPLATE = Path(r"C:\Users\Rodrigo.DESKTOP-2MLHEH3\Desktop\Entregables\VULNERABILIDADES.xlsx")
SOURCE = ROOT / "outputs" / "vulnerabilidades_siaj" / "SIAJ_Plan_Vulnerabilidades_2026_2028.xlsx"
OUTPUT_DIR = ROOT / "outputs" / "vulnerabilidades_siaj"
OUTPUT = OUTPUT_DIR / "SIAJ_Vulnerabilidades_Formato_Seguridad_2026-08-18.xlsx"

SEVERITY_ORDER = ["CRITICAL", "HIGH", "MEDIUM", "LOW"]
STATUS_ORDER = ["RESUELTO", "EN PROCESO", "PENDIENTE"]


def normalize(value):
    return str(value or "").strip().upper()


def copy_cell_style(src, dst):
    if src.has_style:
        dst._style = copy(src._style)
    if src.number_format:
        dst.number_format = src.number_format
    if src.alignment:
        dst.alignment = copy(src.alignment)
    if src.protection:
        dst.protection = copy(src.protection)


def copy_row_style(ws, source_row, target_row, max_col=7):
    for col in range(1, max_col + 1):
        copy_cell_style(ws.cell(source_row, col), ws.cell(target_row, col))
    ws.row_dimensions[target_row].height = ws.row_dimensions[source_row].height


def clear_data(ws):
    max_row = max(ws.max_row, 5)
    for row in range(5, max_row + 1):
        for col in range(1, 6):
            ws.cell(row, col).value = None


def ensure_sheet(wb, severity):
    base_name = "HIGH" if "HIGH" in wb.sheetnames else wb.sheetnames[0]
    if severity in wb.sheetnames:
        ws = wb[severity]
        mapping = header_map(ws)
        if severity == "CRITICAL" and {"vulnerability", "package", "status"} - set(mapping):
            index = wb.sheetnames.index(severity)
            wb.remove(ws)
            ws = wb.copy_worksheet(wb[base_name])
            ws.title = severity
            wb._sheets.remove(ws)
            wb._sheets.insert(index, ws)
        return ws

    source = wb[base_name]
    ws = wb.copy_worksheet(source)
    ws.title = severity
    ws["B2"] = f"VULNERABILIDAD {severity}"
    return ws


def header_map(ws):
    result = {}
    for col in range(1, min(ws.max_column, 7) + 1):
        value = normalize(ws.cell(4, col).value)
        if "VULNERABILIDAD" in value:
            result["vulnerability"] = col
        elif "PACKAGE" in value:
            result["package"] = col
        elif "CATEGORIA" in value or "CATEGOR" in value:
            result["category"] = col
        elif "PAG" in value:
            result["page"] = col
        elif "ESTADO" in value:
            result["status"] = col
    return result


def set_status_legend(ws):
    legend_col = 7
    for col in range(1, ws.max_column + 1):
        if normalize(ws.cell(3, col).value) == "ESTATUS":
            legend_col = col
            break
    ws.cell(3, legend_col).value = "ESTATUS"
    for index, status in enumerate(STATUS_ORDER, start=4):
        ws.cell(index, legend_col).value = status


def read_source_rows():
    wb = openpyxl.load_workbook(SOURCE, data_only=True)
    ws = wb["COMPENDIO"]
    headers = [ws.cell(4, col).value for col in range(1, ws.max_column + 1)]
    indexes = {normalize(name): i + 1 for i, name in enumerate(headers)}

    def col(*names):
        for name in names:
            key = normalize(name)
            if key in indexes:
                return indexes[key]
        raise KeyError(names)

    cols = {
        "severity": col("GRAVEDAD"),
        "vulnerability": col("VULNERABILIDAD"),
        "package": col("PACKAGE", "PACKAGE NAME"),
        "category": col("CATEGORIA / CWE", "CATEGORIA", "CATEGORY"),
        "page": col("PAGINA", "PÁGINA", "PAGINADF"),
        "status": col("ESTATUS", "ESTADO"),
    }

    rows = []
    for row in range(5, ws.max_row + 1):
        severity = normalize(ws.cell(row, cols["severity"]).value)
        if not severity:
            continue
        status = normalize(ws.cell(row, cols["status"]).value) or "PENDIENTE"
        rows.append({
            "severity": severity,
            "vulnerability": ws.cell(row, cols["vulnerability"]).value,
            "package": ws.cell(row, cols["package"]).value,
            "category": ws.cell(row, cols["category"]).value,
            "page": ws.cell(row, cols["page"]).value,
            "status": status,
        })
    return rows


def main():
    OUTPUT_DIR.mkdir(parents=True, exist_ok=True)
    template_wb = openpyxl.load_workbook(TEMPLATE)
    source_rows = read_source_rows()

    by_severity = {severity: [] for severity in SEVERITY_ORDER}
    for row in source_rows:
        by_severity.setdefault(row["severity"], []).append(row)

    for severity in SEVERITY_ORDER:
        ws = ensure_sheet(template_wb, severity)
        ws["B2"] = f"VULNERABILIDAD {severity}"
        clear_data(ws)
        set_status_legend(ws)

        mapping = header_map(ws)
        data = by_severity.get(severity, [])
        if not data:
            data = [{
                "vulnerability": "Sin vulnerabilidades identificadas en esta severidad",
                "package": "",
                "category": "",
                "page": "",
                "status": "RESUELTO",
            }]

        for offset, item in enumerate(data, start=5):
            if offset > ws.max_row:
                ws.append([None] * max(ws.max_column, 7))
            copy_row_style(ws, 5, offset, max_col=7)
            if "vulnerability" in mapping:
                ws.cell(offset, mapping["vulnerability"]).value = item["vulnerability"]
            if "package" in mapping:
                ws.cell(offset, mapping["package"]).value = item["package"]
            if "category" in mapping:
                ws.cell(offset, mapping["category"]).value = item["category"]
            if "page" in mapping:
                ws.cell(offset, mapping["page"]).value = item["page"]
            if "status" in mapping:
                ws.cell(offset, mapping["status"]).value = item["status"]

        if ws.max_row > len(data) + 4:
            ws.delete_rows(len(data) + 5, ws.max_row - (len(data) + 4))

        for table in ws.tables.values():
            table.ref = f"A1:G{max(len(data) + 4, 6)}"

        for col in range(1, min(ws.max_column, 7) + 1):
            ws.column_dimensions[openpyxl.utils.get_column_letter(col)].width = max(
                ws.column_dimensions[openpyxl.utils.get_column_letter(col)].width or 10,
                14,
            )

    ordered = [name for name in SEVERITY_ORDER if name in template_wb.sheetnames]
    ordered.extend([name for name in template_wb.sheetnames if name not in ordered])
    template_wb._sheets = [template_wb[name] for name in ordered]

    template_wb.save(OUTPUT)
    print(OUTPUT)
    for severity in ordered:
        count = len(by_severity.get(severity, []))
        print(f"{severity}: {count}")


if __name__ == "__main__":
    main()
