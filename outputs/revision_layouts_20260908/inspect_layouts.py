from pathlib import Path
import openpyxl, json
root=Path('Layouts_Carga_Masiva_SIAJ')
out=[]
for p in sorted(root.glob('Layout_*.xlsx')):
 w=openpyxl.load_workbook(p,data_only=False)
 out.append('\nFILE '+p.name)
 for s in w:
  out.append(f'SHEET {s.title} rows={s.max_row} cols={s.max_column}')
  for row in s:
   vals=['%s=%s'%(c.column_letter,c.value) for c in row if c.value is not None]
   if vals: out.append(str(row[0].row)+': '+' | '.join(vals))
  out.append('VALIDATIONS '+str([(str(d.sqref),d.type,d.formula1,d.showErrorMessage) for d in s.data_validations.dataValidation]))
Path('outputs/revision_layouts_20260908/excel_inspection.txt').write_text('\n'.join(out),encoding='utf-8')
print('\n'.join(out))
